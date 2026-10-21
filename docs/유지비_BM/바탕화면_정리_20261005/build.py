from pathlib import Path
import json, hashlib, sys
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path(__file__).resolve().parents[3]
BASE=Path(__file__).parent
OUT=BASE/'final'
OUT.mkdir(exist_ok=True)
data=json.loads((BASE/'content.json').read_text(encoding='utf-8'))
FONT='Malgun Gothic'

def font(run,size=None,bold=None):
    run.font.name=FONT
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
    if size: run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    run.font.color.rgb=RGBColor.from_string('000000')

def init():
    d=Document()
    s=d.sections[0]
    s.page_width=Cm(21); s.page_height=Cm(29.7)
    s.top_margin=Cm(1.7); s.bottom_margin=Cm(1.7)
    s.left_margin=Cm(1.8); s.right_margin=Cm(1.8)
    s.footer_distance=Cm(.7)
    for grid in s._sectPr.findall(qn('w:docGrid')): s._sectPr.remove(grid)
    for node in list(d.styles.element.iter(qn('w:pBdr'))): node.getparent().remove(node)
    for name,size in [('Normal',10.5),('Title',23),('Heading 1',13),('Heading 2',11),('Subtitle',10)]:
        st=d.styles[name]
        st.font.name=FONT; st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
        for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            st.element.get_or_add_rPr().rFonts.attrib.pop(qn('w:'+attr),None)
        snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');st.element.get_or_add_pPr().append(snap)
        st.paragraph_format.space_before=Pt(0)
        st.paragraph_format.space_after=Pt(6)
        st.paragraph_format.line_spacing=1.15
        if name in ('Title','Heading 1','Heading 2'): st.font.bold=True
    d.styles['Title'].paragraph_format.space_after=Pt(9)
    d.styles['Heading 1'].paragraph_format.space_before=Pt(9)
    d.core_properties.author='SYNK'
    d.core_properties.subject='2026-10-05 사업모델 정리'
    d.core_properties.comments=''
    return d

def p(d,text,style=None):
    a=d.add_paragraph(style=style); font(a.add_run(text))
    a.paragraph_format.widow_control=True
    return a

def title(d,text):
    p(d,text,'Title')
    a=p(d,'SYNK  /  2026년 10월 5일 기준')
    a.paragraph_format.space_after=Pt(11)
    for r in a.runs: font(r,9)

def table(d,headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers))
    t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Cm(w)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+side); el.set(qn('w:val'),'single');el.set(qn('w:sz'),'5');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    props.append(borders)
    for ri,vals in enumerate([headers]+rows):
        row=t.rows[0] if ri==0 else t.add_row()
        trp=row._tr.get_or_add_trPr()
        nr=OxmlElement('w:cantSplit'); trp.append(nr)
        if ri==0: trp.append(OxmlElement('w:tblHeader'))
        for ci,(cell,text) in enumerate(zip(row.cells,vals)):
            cell.width=Cm(widths[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcp=cell._tc.get_or_add_tcPr()
            mar=OxmlElement('w:tcMar')
            for side,v in [('top',95),('bottom',95),('left',115),('right',115)]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),str(v));e.set(qn('w:type'),'dxa');mar.append(e)
            tcp.append(mar)
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E8EDF6' if ri==0 else ('F8FAFC' if ri%2==0 else 'FFFFFF'));tcp.append(sh)
            a=cell.paragraphs[0]; a.paragraph_format.space_after=Pt(0);a.paragraph_format.line_spacing=1.12
            a.paragraph_format.keep_with_next=(ri==0)
            if ci==0: a.alignment=WD_ALIGN_PARAGRAPH.CENTER
            font(a.add_run(text),9.5,ri==0 or ci==0)
    a=d.add_paragraph();a.paragraph_format.space_after=Pt(0);a.paragraph_format.space_before=Pt(0);a.paragraph_format.line_spacing=1
    font(a.add_run(''),2)

def link(a,label,target):
    rel=a.part.relate_to(target,RT.HYPERLINK,is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rel)
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'244A9B');rp.append(color)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'17');rp.append(size)
    ft=OxmlElement('w:rFonts'); ft.set(qn('w:ascii'),FONT); ft.set(qn('w:hAnsi'),FONT);ft.set(qn('w:eastAsia'),FONT);rp.append(ft)
    r.append(rp);tx=OxmlElement('w:t');tx.text=label;r.append(tx);h.append(r);a._p.append(h)

def source_links(d,ids):
    a=d.add_paragraph();a.paragraph_format.space_before=Pt(7);a.paragraph_format.space_after=Pt(0)
    font(a.add_run('근거  '),8.5)
    for i,id in enumerate(ids):
        v=data['sources'][id]
        target=ROOT/v['path'];assert target.is_file(), target
        if i:font(a.add_run(' · '),8.5)
        link(a,v['name'],target.as_uri())

d=init();summary=data['summary']
d.core_properties.title=summary['title']
title(d,summary['title']);p(d,summary['intro'])
table(d,summary['headers'],summary['rows'],[2.15,7.2,8.05])
p(d,summary['heading'],'Heading 1')
for i,txt in enumerate(summary['points'],1):
    a=p(d,f'{i}. {txt}'); a.paragraph_format.space_after=Pt(5)
    for r in a.runs:font(r,10)
a=p(d,summary['closing'])
for r in a.runs:font(r,8.5)
link(p(d,''),'상세 정리 열기','01_BM_상세.docx')
d.save(OUT/'00_BM_한눈에.docx')

d=init();d.core_properties.title='SYNK 사업모델 상세 정리'
foot=d.sections[0].footer.paragraphs[0]
foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
font(foot.add_run('SYNK  ·  2026.10.05  ·  '),8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
for i,page in enumerate(data['pages']):
    if i:d.add_page_break()
    title(d,page['title']);p(d,page['intro'])
    for block in page['blocks']:
        typ=block['type']
        if typ=='table':
            widths=[3.1,14.3] if len(block['headers'])==2 else [1.1,8.0,8.3]
            table(d,block['headers'],block['rows'],widths)
        elif typ=='heading':p(d,block['text'],'Heading 1')
        elif typ=='p':p(d,block['text'])
        elif typ=='links':source_links(d,block['ids'])
    if page['sources']:source_links(d,page['sources'])
d.save(OUT/'01_BM_상세.docx')

hashes=[]
for id,v in data['sources'].items():
    f=ROOT/v['path']
    hashes.append({'id':id,**v,'sha256':hashlib.sha256(f.read_bytes()).hexdigest()})
(BASE/'source-manifest.json').write_text(json.dumps(hashes,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'outputs':[str(x) for x in OUT.glob('*.docx')],'intended_pages':[1,len(data['pages'])]},ensure_ascii=False))
